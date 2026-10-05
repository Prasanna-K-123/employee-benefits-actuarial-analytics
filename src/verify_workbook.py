"""Check the published workbook and rebuilt formulas against model evidence."""
from pathlib import Path
from math import isclose
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def verify(path):
    cached = load_workbook(path, data_only=True)
    summary = json.loads((ROOT / 'outputs/summary.json').read_text())
    assert isclose(cached['Pension']['I185'].value, summary['pension']['base_liability_proxy'], abs_tol=0.01)
    for cell, key in [('P18','paid_to_date'),('P19','estimated_ultimate'),('P20','ibnr'),('P21','pad'),('P22','reserve_with_pad'),('P29','next_year_cost_proxy')]:
        assert isclose(cached['Health Claims'][cell].value, summary['health'][key], abs_tol=0.01), key
    for ws in cached:
        for row in ws:
            for cell in row:
                assert cell.data_type != 'e', (ws.title, cell.coordinate, cell.value)

def compare_formulas(published, rebuilt):
    left, right = load_workbook(published), load_workbook(rebuilt)
    assert left.sheetnames == right.sheetnames
    for a, b in zip(left, right):
        for row in b:
            for cell in row:
                x, y = a[cell.coordinate].value, cell.value
                if x in (None, '') and y in (None, ''):
                    continue
                if isinstance(x, (int,float)) and isinstance(y,(int,float)):
                    assert isclose(x,y,abs_tol=1e-8), (a.title,cell.coordinate)
                else:
                    assert x == y, (a.title,cell.coordinate,x,y)
    assert right['Pension']['E4'].number_format == '0'

if __name__ == '__main__':
    import sys
    if len(sys.argv) == 3:
        compare_formulas(*sys.argv[1:])
    else:
        verify(sys.argv[1])
    print('Workbook evidence verified')
