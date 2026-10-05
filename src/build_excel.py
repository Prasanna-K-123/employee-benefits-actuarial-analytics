from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUTPUT = ROOT / "Employee_Benefits_Actuarial_Model.xlsx"


def _title(ws, text: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, text)
    c.font = Font(bold=True, color="FFFFFF", size=12)
    c.fill = PatternFill("solid", fgColor="17365D")
    c.alignment = Alignment(horizontal="center")


def _header(ws, row: int, start_col: int, labels: list[str]) -> None:
    for j, label in enumerate(labels, start=start_col):
        c = ws.cell(row, j, label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def build() -> Path:
    pension = pd.read_csv(DATA / "synthetic_pension_population.csv")
    health = pd.read_csv(DATA / "synthetic_health_claims_triangle.csv")

    wb = Workbook()
    ass = wb.active
    ass.title = "Assumptions"
    _title(ass, "EMPLOYEE BENEFITS ACTUARIAL MODEL - EDUCATIONAL / SYNTHETIC", 4)
    _header(ass, 3, 1, ["Pension assumptions", "Value", "Health assumptions", "Value"])
    inputs = [
        ("Discount rate", 0.05, "PAD / risk margin", 0.05),
        ("Salary growth", 0.04, "Annual claim trend", 0.07),
        ("Accrual rate", 0.015, "", ""),
        ("Retirement age", 65, "", ""),
        ("Annuity discount rate", 0.04, "", ""),
        ("Annuity years", 15, "", ""),
    ]
    for r, values in enumerate(inputs, 4):
        for c, v in enumerate(values, 1):
            ass.cell(r, c, v)
    for cell in ("B4", "B5", "B6", "B8", "D4", "D5"):
        ass[cell].number_format = "0.0%"
    ass["A11"] = "MODEL BOUNDARY"
    ass["A11"].font = Font(bold=True)
    ass["A12"] = "All employee and claims data are synthetic."
    ass["A13"] = "Retirement output is an illustrative liability proxy, not an ASC 715 / IAS 19 valuation."
    ass["A14"] = "Health output is an educational chain-ladder IBNR / PAD illustration, not a booked reserve."

    pen = wb.create_sheet("Pension")
    _title(pen, "RETIREMENT / DEFINED-BENEFIT LIABILITY PROXY", 9)
    labels = [
        "Employee ID", "Age", "Service Years", "Salary ($)", "Years to Retirement",
        "Projected Final Salary ($)", "Accrued Annual Benefit ($)",
        "PV at Retirement ($)", "Liability Proxy ($)"
    ]
    _header(pen, 3, 1, labels)
    for i, row in pension.iterrows():
        r = 4 + i
        pen.cell(r, 1, row["employee_id"])
        pen.cell(r, 2, int(row["age"]))
        pen.cell(r, 3, int(row["service_years"]))
        pen.cell(r, 4, float(row["salary"]))
        pen.cell(r, 5, f"=MAX(Assumptions!$B$7-B{r},0)")
        pen.cell(r, 6, f"=D{r}*(1+Assumptions!$B$5)^E{r}")
        pen.cell(r, 7, f"=Assumptions!$B$6*F{r}*C{r}")
        pen.cell(r, 8, f"=G{r}*IF(Assumptions!$B$8=0,Assumptions!$B$9,(1-(1+Assumptions!$B$8)^(-Assumptions!$B$9))/Assumptions!$B$8)")
        pen.cell(r, 9, f"=H{r}/(1+Assumptions!$B$4)^E{r}")
    for col in range(4, 10):
        for r in range(4, 4 + len(pension)):
            pen.cell(r, col).number_format = '$#,##0'
    pen.freeze_panes = "A4"
    for r in range(4, 4 + len(pension)):
        pen.cell(r, 5).number_format = "0"
    total_row = 5 + len(pension)
    pen.cell(total_row, 8, "Total liability proxy ($)").font = Font(bold=True)
    pen.cell(total_row, 9, f"=SUM(I4:I{3 + len(pension)})").number_format = '$#,##0'
    pen.cell(total_row, 9).font = Font(bold=True)

    hea = wb.create_sheet("Health Claims")
    _title(hea, "HEALTH & WELFARE CLAIMS DEVELOPMENT / IBNR", 18)
    triangle_cols = [c for c in health.columns if c.startswith("dev_")]
    _header(hea, 3, 1, ["Accident Month"] + triangle_cols)
    for i, row in health.iterrows():
        r = 4 + i
        hea.cell(r, 1, row["accident_month"])
        for j, col in enumerate(triangle_cols, 2):
            v = row[col]
            if pd.notna(v):
                hea.cell(r, j, float(v)).number_format = '$#,##0'
    hea["A17"] = "Age-to-age factor"
    for j in range(2, 13):
        current = get_column_letter(j)
        nxt = get_column_letter(j + 1)
        last = 16 - j
        hea.cell(17, j, f"=SUM({nxt}4:{nxt}{last})/SUM({current}4:{current}{last})")
        hea.cell(17, j).number_format = "0.0000x"
    hea["M17"] = 1.0
    hea["A18"] = "CDF to ultimate"
    hea["M18"] = 1.0
    for j in range(12, 1, -1):
        current = get_column_letter(j)
        nxt = get_column_letter(j + 1)
        hea.cell(18, j, f"={current}17*{nxt}18")
        hea.cell(18, j).number_format = "0.0000x"

    _header(hea, 3, 15, ["Accident Month", "Latest Paid ($)", "Estimated Ultimate ($)", "IBNR ($)"])
    for i in range(12):
        r = 4 + i
        latest_col = get_column_letter(13 - i)
        hea.cell(r, 15, f"M{i+1:02d}")
        hea.cell(r, 16, f"={latest_col}{r}")
        hea.cell(r, 17, f"=P{r}*{latest_col}$18")
        hea.cell(r, 18, f"=Q{r}-P{r}")
        for c in range(16, 19):
            hea.cell(r, c).number_format = '$#,##0'
    hea["O17"] = "Summary"
    hea["P17"] = "Value"
    hea["O18"] = "Paid to date"; hea["P18"] = "=SUM(P4:P15)"
    hea["O19"] = "Estimated ultimate"; hea["P19"] = "=SUM(Q4:Q15)"
    hea["O20"] = "IBNR"; hea["P20"] = "=SUM(R4:R15)"
    hea["O21"] = "PAD / risk margin"; hea["P21"] = "=P20*Assumptions!$D$4"
    hea["O22"] = "Reserve incl. PAD"; hea["P22"] = "=P20+P21"
    for r in range(18, 23):
        hea.cell(r, 16).number_format = '$#,##0'
    hea["O24"] = "Claim-trend sensitivity"
    hea["P24"] = "Next-year cost proxy ($)"
    for r, trend in zip(range(25, 28), (0.05, 0.07, 0.09)):
        hea.cell(r, 15, f"{trend:.0%}")
        hea.cell(r, 16, f"=P19*(1+{trend})").number_format = '$#,##0'
    hea.freeze_panes = "A4"
    hea["O29"] = "Selected trend cost proxy"
    hea["P29"] = "=P19*(1+Assumptions!$D$5)"
    hea["P29"].number_format = '$#,##0'

    for ws in wb.worksheets:
        for column in ws.columns:
            width = min(max((len(str(c.value)) if c.value is not None else 0) for c in column) + 2, 26)
            ws.column_dimensions[get_column_letter(column[0].column)].width = max(width, 10)

    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
